# -*- coding: utf-8 -*-
"""md2docx.py —— 把 Markdown 方法规范转成 Word 文档（中文正确显示）

要点：中文必须显式设置 w:eastAsia 字体，否则 Word 里可能显示成方框或回退字体。
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = sys.argv[1]
DST = sys.argv[2]

doc = Document()

# ---- 设置正文默认字体（中英文都要设）----
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.25


def set_cjk(run, font="微软雅黑"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def add_heading(text, level):
    sizes = {0: 18, 1: 15, 2: 13, 3: 11.5}
    colors = {0: RGBColor(0x1F, 0x38, 0x64), 1: RGBColor(0x1F, 0x38, 0x64),
              2: RGBColor(0x2E, 0x5C, 0x9A), 3: RGBColor(0x40, 0x40, 0x40)}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level <= 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(sizes.get(level, 11))
    r.font.color.rgb = colors.get(level, RGBColor(0, 0, 0))
    set_cjk(r)
    return p


def add_code(lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("\n".join(lines))
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    # 处理行内粗体 **xx**
    for i, seg in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2]); r.bold = True
        else:
            r = p.add_run(seg)
        set_cjk(r)
    return p


def add_table(rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            val = row[ci] if ci < len(row) else ""
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            # 表格里也用行内粗体
            for seg in re.split(r"(\*\*[^*]+\*\*)", val):
                if not seg:
                    continue
                if seg.startswith("**") and seg.endswith("**"):
                    r = p.add_run(seg[2:-2]); r.bold = True
                else:
                    r = p.add_run(seg)
                r.font.size = Pt(9.5)
                if ri == 0:
                    r.bold = True
                set_cjk(r)
    doc.add_paragraph()
    return t


lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    line = lines[i].rstrip()

    # 代码块
    if line.startswith("```"):
        buf = []
        i += 1
        while i < len(lines) and not lines[i].startswith("```"):
            buf.append(lines[i]); i += 1
        add_code(buf)
        i += 1
        continue

    # 表格
    if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|", lines[i + 1]):
        rows = []
        header = [c.strip() for c in line.strip("|").split("|")]
        rows.append(header)
        i += 2
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        add_table(rows)
        continue

    # 标题
    m = re.match(r"^(#{1,4})\s+(.*)", line)
    if m:
        add_heading(m.group(2).strip(), len(m.group(1)) - 1)
        i += 1
        continue

    # 引用块（同样要处理行内粗体，否则 Word 里会留下 ** 星号）
    if line.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        for seg in re.split(r"(\*\*[^*]+\*\*)", line.lstrip("> ").strip()):
            if not seg:
                continue
            if seg.startswith("**") and seg.endswith("**"):
                r = p.add_run(seg[2:-2]); r.bold = True
            else:
                r = p.add_run(seg)
            r.italic = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            set_cjk(r)
        i += 1
        continue

    # 分隔线
    if re.match(r"^-{3,}$", line.strip()):
        p = doc.add_paragraph()
        r = p.add_run("─" * 40)
        r.font.color.rgb = RGBColor(0xC0, 0xC0, 0xC0)
        set_cjk(r)
        i += 1
        continue

    # 列表
    m = re.match(r"^(\s*)[-*]\s+(.*)", line)
    if m:
        add_bullet(m.group(2), 1 if len(m.group(1)) >= 2 else 0)
        i += 1
        continue

    # 有序列表
    m = re.match(r"^(\s*)(\d+)\.\s+(.*)", line)
    if m:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        for seg in re.split(r"(\*\*[^*]+\*\*)", m.group(3)):
            if not seg:
                continue
            if seg.startswith("**") and seg.endswith("**"):
                r = p.add_run(seg[2:-2]); r.bold = True
            else:
                r = p.add_run(seg)
            set_cjk(r)
        i += 1
        continue

    # 空行
    if not line.strip():
        i += 1
        continue

    # 普通段落
    p = doc.add_paragraph()
    for seg in re.split(r"(\*\*[^*]+\*\*)", line):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2]); r.bold = True
        else:
            r = p.add_run(seg)
        set_cjk(r)
    i += 1

doc.save(DST)
print("已生成:", DST)
print("段落数:", len(doc.paragraphs), " 表格数:", len(doc.tables))
